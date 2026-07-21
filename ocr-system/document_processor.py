# document_processor.py
import os
from PIL import Image, ImageEnhance, ImageFilter
import PyPDF2
from docx import Document as DocxDocument
import pytesseract
from typing import List, Optional
import logging
from config import config

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Handles document processing, OCR, and text extraction"""

    def __init__(self):
        # Configure Tesseract path if provided
        if config.TESSERACT_PATH and os.path.exists(config.TESSERACT_PATH):
            pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_PATH

        self.supported_extensions = config.ALLOWED_EXTENSIONS
        logger.info(f"DocumentProcessor initialized with {len(self.supported_extensions)} supported file types")

    def process_document(self, file_path: str) -> str:
        """Process document and extract text based on file type"""
        ext = os.path.splitext(file_path)[1].lower()

        if ext not in self.supported_extensions:
            raise ValueError(f"Unsupported file type: {ext}. Allowed: {', '.join(self.supported_extensions)}")

        try:
            if ext == '.pdf':
                return self._process_pdf(file_path)
            elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp']:
                return self._process_image(file_path)
            elif ext == '.docx':
                return self._process_docx(file_path)
            elif ext == '.txt':
                return self._process_txt(file_path)
            else:
                raise ValueError(f"Unsupported file type: {ext}")

        except Exception as e:
            logger.error(f"Error processing {file_path}: {str(e)}")
            return f"[Error processing document: {str(e)}]"

    def _process_pdf(self, file_path: str) -> str:
        """Extract text from PDF"""
        text = ""
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"

            # If no text found, try OCR
            if not text.strip():
                text = self._ocr_pdf_images(file_path)

        except Exception as e:
            logger.error(f"PDF processing error: {e}")
            text = f"[PDF processing error: {str(e)}]"

        return text.strip()

    def _ocr_pdf_images(self, file_path: str) -> str:
        """Convert PDF to images and OCR (requires pdf2image)"""
        try:
            from pdf2image import convert_from_path
            text = ""
            images = convert_from_path(file_path, dpi=200)

            for i, image in enumerate(images):
                page_text = self._ocr_image(image)
                if page_text:
                    text += f"--- Page {i+1} ---\n{page_text}\n\n"

            return text
        except Exception as e:
            logger.error(f"PDF OCR error: {e}")
            return f"[PDF OCR failed: {str(e)}]"

    def _process_image(self, file_path: str) -> str:
        """Extract text from image using PIL and Tesseract"""
        try:
            image = Image.open(file_path)
            return self._ocr_image(image)
        except Exception as e:
            logger.error(f"Image processing error: {e}")
            return f"[Image OCR failed: {str(e)}]"

    def _ocr_image(self, image) -> str:
        """Perform OCR on PIL image with preprocessing"""
        try:
            # Convert to grayscale
            if image.mode != 'L':
                gray = image.convert('L')
            else:
                gray = image

            # Enhance contrast
            enhancer = ImageEnhance.Contrast(gray)
            enhanced = enhancer.enhance(2.0)

            # Sharpen
            sharpened = enhanced.filter(ImageFilter.SHARPEN)

            # Resize if too large (for better OCR)
            max_size = 2000
            if sharpened.width > max_size or sharpened.height > max_size:
                sharpened.thumbnail((max_size, max_size), Image.Resampling.LANCZOS)

            # OCR
            text = pytesseract.image_to_string(sharpened, lang='eng')
            return text.strip()

        except Exception as e:
            logger.error(f"OCR error: {e}")
            return f"[OCR error: {str(e)}]"

    def _process_docx(self, file_path: str) -> str:
        """Extract text from DOCX"""
        try:
            doc = DocxDocument(file_path)
            text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text.append(cell.text)

            return "\n".join(text)

        except Exception as e:
            logger.error(f"DOCX processing error: {e}")
            return f"[DOCX processing error: {str(e)}]"

    def _process_txt(self, file_path: str) -> str:
        """Read text file with encoding fallback"""
        encodings = ['utf-8', 'latin-1', 'cp1252', 'utf-16']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue

        # Fallback: read as binary and decode with errors='replace'
        with open(file_path, 'rb') as f:
            content = f.read()
            return content.decode('utf-8', errors='replace')

    def chunk_text(self, text: str, chunk_size: Optional[int] = None) -> List[str]:
        """Split text into chunks for better search"""
        if not text:
            return []

        chunk_size = chunk_size or config.CHUNK_SIZE
        words = text.split()
        chunks = []
        current_chunk = []
        current_size = 0

        for word in words:
            current_chunk.append(word)
            current_size += len(word) + 1

            if current_size >= chunk_size:
                chunks.append(' '.join(current_chunk))
                current_chunk = []
                current_size = 0

        if current_chunk:
            chunks.append(' '.join(current_chunk))

        logger.info(f"Split text into {len(chunks)} chunks")
        return chunks

    def get_page_count(self, file_path: str) -> int:
        """Get page count for document"""
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == '.pdf':
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    return len(reader.pages)
            elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp']:
                return 1
            else:
                return 1
        except:
            return 0
